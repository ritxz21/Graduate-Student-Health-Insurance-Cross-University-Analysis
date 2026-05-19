"""
pipeline.py — Main orchestration script for the grad insurance extraction pipeline.

Usage:
    # From the project root (with venv activated):
    python scripts/pipeline.py                          # run all universities
    python scripts/pipeline.py --ids harvard stanford   # run specific schools
    python scripts/pipeline.py --limit 5               # run first 5 schools
    python scripts/pipeline.py --dry-run               # print plan, don't fetch

What it does per university:
  1. Fetch HTML from source_url_primary (+ secondary if present)
  2. Discover and download PDF links into data/raw/<university_id>/
  3. Extract text from all PDFs
  4. Run field_extractors.extract_all_fields() on combined text
  5. Write row to data/extracted_raw.csv
  6. Update the matching row in data/tracking_sheet.xlsx

Logging is written to outputs/pipeline.log.
"""

from __future__ import annotations

import argparse
import csv
import datetime
import logging
import sys
import time
from pathlib import Path

# Add scripts/ to path so we can import sibling modules
sys.path.insert(0, str(Path(__file__).resolve().parent))

from config import (
    UNIVERSITIES, TRACKING_FIELDS, SYSTEM_PLAN_MAP,
    DATA_DIR, RAW_PDF_DIR, OUTPUTS_DIR, TRACKING_XLSX,
    REQUEST_DELAY_SECONDS,
)
from html_scraper import scrape_university, download_pdf
from pdf_extractor import extract_text_from_pdf
from field_extractors import extract_all_fields

# ---------------------------------------------------------------------------
# Logging setup
# ---------------------------------------------------------------------------
log_path = OUTPUTS_DIR / "pipeline.log"
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    handlers=[
        logging.FileHandler(log_path, encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# CSV output helpers
# ---------------------------------------------------------------------------
EXTRACTED_CSV = DATA_DIR / "extracted_raw.csv"


def _csv_header() -> list[str]:
    return TRACKING_FIELDS + ["html_text_chars", "pdf_files_downloaded", "pipeline_error"]


def _open_csv_writer():
    """Open (or append to) the raw extraction CSV and return (file_handle, writer)."""
    write_header = not EXTRACTED_CSV.exists()
    fh = open(EXTRACTED_CSV, "a", newline="", encoding="utf-8")
    writer = csv.DictWriter(fh, fieldnames=_csv_header(), extrasaction="ignore")
    if write_header:
        writer.writeheader()
    return fh, writer


# ---------------------------------------------------------------------------
# Tracking sheet helpers
# ---------------------------------------------------------------------------

def _load_tracking_sheet():
    """Return openpyxl workbook and the Tracking worksheet."""
    import openpyxl
    wb = openpyxl.load_workbook(str(TRACKING_XLSX))
    ws = wb["Tracking"]
    return wb, ws


def _find_tracking_row(ws, university_id: str, plan_variant: str = "default") -> int | None:
    """
    Find the 1-based row number in the Tracking sheet for the given (uni_id, plan_variant).
    Returns None if not found.
    """
    header_row = None
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row[0] == "university_id":
            header_row = row_idx
        if header_row and row[0] == university_id and row[1] == plan_variant:
            return row_idx
    return None


def _get_col_index(ws) -> dict[str, int]:
    """Return {field_name: col_index (0-based)} from the header row."""
    for row in ws.iter_rows(values_only=True):
        if row[0] == "university_id":
            return {cell: idx for idx, cell in enumerate(row) if cell}
    return {}


def _update_tracking_row(ws, row_idx: int, col_index: dict[str, int], data: dict) -> None:
    """Write data dict values into the matching columns of row_idx."""
    for field, value in data.items():
        if field in col_index:
            col = col_index[field] + 1  # openpyxl is 1-based
            ws.cell(row=row_idx, column=col, value=value)


# ---------------------------------------------------------------------------
# Per-university pipeline
# ---------------------------------------------------------------------------

def _build_empty_row(uni: dict) -> dict:
    """Build a row dict pre-filled with identity fields and defaults."""
    row: dict = {f: None for f in TRACKING_FIELDS}
    row.update({
        "university_id":    uni["university_id"],
        "plan_variant":     "default",
        "university_name":  uni["university_name"],
        "university_type":  uni["university_type"],
        "control":          uni["control"],
        "region":           uni["region"],
        "state":            uni["state"],
        "source_url_primary":   uni["source_url_primary"],
        "source_url_secondary": uni.get("source_url_secondary", ""),
        "system_plan_id":       SYSTEM_PLAN_MAP.get(uni["university_id"], ""),
        "grad_union_present":   0,
        "grad_union_strength":  "none",
        "grad_union_name":      "",
        "collection_status":    "not_started",
        "data_quality_flag":    "unverified",
        "extraction_method":    "manual",
        "manual_override_fields": "",
        "notes":                "",
        # Subsidy defaults — usually need manual verification
        "premium_subsidy_type":         "none",
        "premium_subsidy_eligibility":  "",
    })
    return row


def _assess_quality(row: dict) -> str:
    """Simple quality flag based on how many key fields were populated."""
    key_fields = [
        "premium_annual_usd", "deductible_in_network_usd",
        "oop_max_in_network_usd", "mental_health_covered",
    ]
    populated = sum(1 for f in key_fields if row.get(f) not in (None, "", "NA"))
    if populated == 4:
        return "clean"
    if populated >= 2:
        return "partial"
    return "needs_review"


def process_university(uni: dict, download_pdfs: bool = True) -> dict:
    """
    Full pipeline for one university.

    Returns the completed row dict (ready to write to CSV / tracking sheet).
    """
    uid = uni["university_id"]
    log.info(f"[{uid}] Starting...")

    row = _build_empty_row(uni)
    row["data_collection_date"] = datetime.date.today().isoformat()

    # ── Step 1: Fetch HTML ──────────────────────────────────────────────────
    log.info(f"[{uid}] Fetching HTML...")
    scrape_result = scrape_university(uni)

    html_text = scrape_result["html_text"]
    pdf_urls  = scrape_result["pdf_urls"]
    row["source_type"] = scrape_result["source_type"]

    if scrape_result["error"]:
        log.warning(f"[{uid}] HTML fetch issue: {scrape_result['error']}")

    # ── Step 2: Download + extract PDFs ─────────────────────────────────────
    uni_pdf_dir = RAW_PDF_DIR / uid
    pdf_texts: list[str] = []
    pdfs_downloaded = 0

    if download_pdfs and pdf_urls:
        log.info(f"[{uid}] Found {len(pdf_urls)} PDF link(s). Downloading up to 5...")
        for i, pdf_url in enumerate(pdf_urls[:5]):   # cap at 5 PDFs per school
            pdf_filename = f"{uid}_{i+1:02d}.pdf"
            save_path = uni_pdf_dir / pdf_filename
            time.sleep(0.5)
            ok = download_pdf(pdf_url, str(save_path))
            if ok:
                pdfs_downloaded += 1
                log.info(f"[{uid}]   Downloaded: {pdf_filename}")
                text = extract_text_from_pdf(save_path)
                if text.startswith("__SCANNED_PDF__"):
                    log.warning(f"[{uid}]   {pdf_filename} appears to be scanned image — manual OCR needed")
                elif text.startswith("__FILE_NOT_FOUND__"):
                    pass
                else:
                    pdf_texts.append(text)
            else:
                log.warning(f"[{uid}]   Failed to download: {pdf_url}")

    # ── Step 3: Combine text ────────────────────────────────────────────────
    combined_text = html_text
    if pdf_texts:
        combined_text += "\n\n" + "\n\n".join(pdf_texts)

    # ── Step 4: Extract all fields ──────────────────────────────────────────
    if combined_text.strip():
        log.info(f"[{uid}] Running field extractors on {len(combined_text):,} chars...")
        extracted = extract_all_fields(combined_text, uni_id=uid)
        row.update(extracted)
        row["extraction_method"] = "hybrid" if pdf_texts else "auto_html"
    else:
        log.warning(f"[{uid}] No text to extract from — marking as needs_review")
        row["extraction_method"] = "manual"
        row["notes"] = (row.get("notes") or "") + " No text extracted; manual review needed."

    # ── Step 5: Quality flag ────────────────────────────────────────────────
    row["data_quality_flag"] = _assess_quality(row)
    row["collection_status"] = (
        "extracted" if row["data_quality_flag"] in ("clean", "partial") else "in_progress"
    )

    # ── Diagnostics (not stored in tracking sheet) ──────────────────────────
    row["html_text_chars"]        = len(html_text)
    row["pdf_files_downloaded"]   = pdfs_downloaded

    log.info(
        f"[{uid}] Done. Quality: {row['data_quality_flag']} | "
        f"Premium: {row.get('premium_annual_usd')} | "
        f"Deductible: {row.get('deductible_in_network_usd')} | "
        f"OOP max: {row.get('oop_max_in_network_usd')}"
    )
    return row


# ---------------------------------------------------------------------------
# Batch runner
# ---------------------------------------------------------------------------

def run_pipeline(
    university_ids: list[str] | None = None,
    limit: int | None = None,
    dry_run: bool = False,
    download_pdfs: bool = True,
) -> list[dict]:
    """
    Run the full pipeline for a list of universities.

    Args:
        university_ids: subset to run (None = all)
        limit:          max number to process
        dry_run:        if True, just print the plan without fetching
        download_pdfs:  whether to attempt PDF downloads

    Returns:
        List of completed row dicts.
    """
    targets = UNIVERSITIES
    if university_ids:
        id_set = set(university_ids)
        targets = [u for u in UNIVERSITIES if u["university_id"] in id_set]
    if limit:
        targets = targets[:limit]

    log.info(f"Pipeline starting — {len(targets)} universities to process")

    if dry_run:
        for u in targets:
            print(f"  {u['university_id']:30s}  {u['source_url_primary']}")
        return []

    # Load tracking sheet once
    try:
        wb, ws = _load_tracking_sheet()
        col_index = _get_col_index(ws)
        tracking_available = True
    except Exception as exc:
        log.warning(f"Could not load tracking sheet: {exc}. Will write CSV only.")
        wb = ws = col_index = None
        tracking_available = False

    # Open CSV for appending
    fh, csv_writer = _open_csv_writer()

    results: list[dict] = []
    for i, uni in enumerate(targets):
        if i > 0:
            time.sleep(REQUEST_DELAY_SECONDS)
        try:
            row = process_university(uni, download_pdfs=download_pdfs)
        except Exception as exc:
            log.error(f"[{uni['university_id']}] Unhandled error: {exc}", exc_info=True)
            row = _build_empty_row(uni)
            row["data_quality_flag"] = "needs_review"
            row["notes"] = f"Pipeline error: {exc}"
            row["pipeline_error"] = str(exc)

        # Write to CSV
        csv_writer.writerow(row)
        fh.flush()

        # Update tracking sheet
        if tracking_available:
            try:
                row_idx = _find_tracking_row(ws, uni["university_id"])
                if row_idx:
                    _update_tracking_row(ws, row_idx, col_index, row)
                else:
                    log.warning(f"[{uni['university_id']}] Row not found in tracking sheet — CSV only")
            except Exception as exc:
                log.warning(f"[{uni['university_id']}] Could not update tracking sheet: {exc}")

        results.append(row)

    fh.close()

    # Save tracking sheet
    if tracking_available:
        try:
            wb.save(str(TRACKING_XLSX))
            log.info(f"Tracking sheet saved: {TRACKING_XLSX}")
        except Exception as exc:
            log.error(f"Could not save tracking sheet: {exc}")

    log.info(f"Pipeline complete. Results written to: {EXTRACTED_CSV}")
    return results


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------

def _parse_args():
    parser = argparse.ArgumentParser(
        description="Grad insurance extraction pipeline"
    )
    parser.add_argument(
        "--ids", nargs="+", metavar="UNIVERSITY_ID",
        help="Run only these university IDs (space-separated)"
    )
    parser.add_argument(
        "--limit", type=int, default=None,
        help="Process at most N universities"
    )
    parser.add_argument(
        "--dry-run", action="store_true",
        help="Print what would be run without fetching anything"
    )
    parser.add_argument(
        "--no-pdf", action="store_true",
        help="Skip PDF downloads (faster, HTML-only extraction)"
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = _parse_args()
    run_pipeline(
        university_ids=args.ids,
        limit=args.limit,
        dry_run=args.dry_run,
        download_pdfs=not args.no_pdf,
    )
